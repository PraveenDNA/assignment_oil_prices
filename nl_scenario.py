#!/usr/bin/env python3
"""
Natural Language Stress Scenario Engine — Brent Crude
======================================================

A risk analyst writes a scenario in plain English.
Claude extracts the quantitative model parameters, the GJR-GARCH model runs,
and Claude interprets the numbers back into plain English for a risk committee.

Flow
----
  Analyst writes  →  Claude extracts params  →  GJR-GARCH simulates
  →  Claude interprets results  →  Word doc + console briefing

Usage
-----
    python nl_scenario.py
    python nl_scenario.py --query "What if OPEC floods the market?"
    python nl_scenario.py --query "..." --n-paths 500

Requirements
------------
    ANTHROPIC_API_KEY environment variable (or --api-key flag)
    pip install anthropic

    On AWS this call would go to Amazon Bedrock (Claude) instead of
    the Anthropic API directly — same SDK, different base_url + credentials.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import textwrap
import warnings
from datetime import date
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT / "src"))

# ── model & doc imports ───────────────────────────────────────────────────────
import io
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from brent_stress.data import fetch_brent
from brent_stress.model import fit_gjr_garch

SEED    = 42
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)


# ═══════════════════════════════════════════════════════════════════════════════
# PROMPT TEMPLATES
# ═══════════════════════════════════════════════════════════════════════════════

EXTRACTION_SYSTEM = """\
You are a quantitative risk model assistant at an energy trading firm.
You translate plain-English commodity stress scenarios into precise model parameters
for a GJR-GARCH(1,1) + Student-t stress-test model calibrated on Brent crude (BZ=F).
You respond ONLY with valid JSON — no prose, no markdown, no code fences.\
"""

EXTRACTION_USER = """\
Translate the analyst's scenario into GJR-GARCH model parameters.

PARAMETER GUIDE
───────────────────────────────────────────────────────────────────────────────
calibration_start  ISO date. Which historical regime best matches the scenario?
                   "2016-01-01" → calm / recovering (2016-2018 OPEC compliance)
                   "2019-06-01" → pre-COVID + crash  (includes 2020 demand collapse)
                   "2021-01-01" → Russia-Ukraine era (high vol, supply disruption)
                   "2022-01-01" → crisis regime only (2022-2023 elevated vol)
                   "2015-01-01" → full history baseline

calibration_end    ISO date or null. null = use data up to today.

horizon            Projection in trading days.
                   Acute shock    →  21–63   (1–3 months)
                   Medium term    →  63–126  (3–6 months)
                   Full year      →  252
                   Multi-year     →  504

initial_vol_pct    Starting daily vol (%). null = unconditional fitted value.
                   Normal market  →  null or 1–2 %
                   Elevated       →  3–4 %
                   Crisis onset   →  5–6 %
                   Peak panic     →  7–10 %

gamma_multiplier   Leverage effect multiplier (1.0 = baseline).
                   Symmetric shock           →  1.0
                   Moderately asymmetric     →  1.5
                   Strongly downside-biased  →  2.0–3.0

nu_override        Student-t degrees of freedom. null = use fitted (~5.2).
                   Extreme tail risk   →  3.0–4.0
                   Standard heavy tail →  null
                   Orderly market      →  7.0–10.0

mu_shock_pct       Constant daily drift added to model mean (% per day).
                   Supply shock / price spike  →  +0.03 to +0.10
                   Demand collapse / bearish   →  -0.10 to -0.03
                   No directional view         →  0.0

scenario_name      Short label ≤ 8 words.

reasoning          2–3 sentences justifying each key parameter choice.
───────────────────────────────────────────────────────────────────────────────

Respond with ONLY this JSON (no other text):
{{
  "scenario_name":    "...",
  "calibration_start":"YYYY-MM-DD",
  "calibration_end":   "YYYY-MM-DD or null",
  "horizon":           <int>,
  "initial_vol_pct":   <float or null>,
  "gamma_multiplier":  <float>,
  "nu_override":       <float or null>,
  "mu_shock_pct":      <float>,
  "reasoning":         "..."
}}

Analyst scenario: "{query}"\
"""

INTERPRETATION_SYSTEM = """\
You are a senior quantitative risk analyst presenting stress-test results to a risk committee.
Your writing is precise, jargon-appropriate for a risk committee audience, and actionable.
You never restate numbers verbatim — you interpret what they mean.\
"""

INTERPRETATION_USER = """\
Write a three-paragraph risk committee briefing note for the following scenario and results.

SCENARIO: {scenario_name}
ANALYST QUERY: "{query}"

MODEL: GJR-GARCH(1,1) + Student-t on Brent crude (BZ=F)
CALIBRATION: {calibration_start} to {calibration_end}
HORIZON: {horizon} trading days
PATHS: {n_paths:,} Monte Carlo simulations

RISK METRICS
────────────────────────────────────────
Annualised volatility      {ann_vol:.1f} %
Daily VaR 95%              {var95:.2f} %
Daily VaR 99%              {var99:.2f} %
Daily ES 99%               {es99:.2f} %
────────────────────────────────────────
Horizon cumulative  P1     {cum_p01:.1f} %
Horizon cumulative  P10    {cum_p10:.1f} %
Horizon cumulative median  {cum_p50:.1f} %
Horizon cumulative  P90    {cum_p90:.1f} %
────────────────────────────────────────
Mean max drawdown          {mean_dd:.1f} %
Worst drawdown (P95)       {p95_dd:.1f} %
Prob(drawdown > 20 %)      {prob_dd_gt20:.1f} %
Prob(drawdown > 40 %)      {prob_dd_gt40:.1f} %
Prob(total loss > 30 %)    {prob_loss_gt30:.1f} %
────────────────────────────────────────

PARAGRAPH STRUCTURE:

§1 HEADLINE RISK — What the numbers say in plain English.
   Translate VaR and drawdown into intuitive terms (e.g. "on the worst one-in-a-hundred
   trading days, Brent loses X %, roughly equivalent to $Y/barrel on a $Z current price").
   State the realistic range of outcomes across the horizon.

§2 DRIVERS — Why this scenario produces these results.
   Link model parameters to economic mechanism. Explain how initial vol, tail heaviness,
   and leverage amplification combine to shape the loss distribution in this specific context.

§3 RISK MANAGEMENT ACTIONS — What a risk manager should do.
   Name specific instruments (e.g. put options, calendar spreads, collar strategies,
   physical storage hedges). State concrete trigger levels for position limit review.
   Be actionable — a committee should leave with three specific decisions to make.\
"""


# ═══════════════════════════════════════════════════════════════════════════════
# CLAUDE CALLS
# ═══════════════════════════════════════════════════════════════════════════════

def _get_client(api_key: str | None):
    import anthropic
    key = api_key or os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        raise SystemExit(
            "\n  ✗  ANTHROPIC_API_KEY not set.\n"
            "     Export it:  export ANTHROPIC_API_KEY=sk-ant-...\n"
            "     Or pass:    python nl_scenario.py --api-key sk-ant-...\n"
        )
    return anthropic.Anthropic(api_key=key)


def extract_params(query: str, client) -> dict:
    """
    Call Claude to translate a natural-language scenario into model parameters.
    Returns a validated dict ready to pass to the simulation engine.
    """
    response = client.messages.create(
        model="claude-opus-4-6",
        max_tokens=512,
        system=EXTRACTION_SYSTEM,
        messages=[{"role": "user", "content": EXTRACTION_USER.format(query=query)}],
    )
    raw = response.content[0].text.strip()

    # Strip any accidental markdown fences
    if raw.startswith("```"):
        raw = "\n".join(raw.split("\n")[1:])
    if raw.endswith("```"):
        raw = "\n".join(raw.split("\n")[:-1])

    params = json.loads(raw)

    # Coerce types and fill defaults
    params["horizon"]          = int(params.get("horizon", 252))
    params["gamma_multiplier"] = float(params.get("gamma_multiplier", 1.0))
    params["mu_shock_pct"]     = float(params.get("mu_shock_pct", 0.0))
    params["calibration_end"]  = params.get("calibration_end") or None
    params["initial_vol_pct"]  = (
        float(params["initial_vol_pct"]) if params.get("initial_vol_pct") else None
    )
    params["nu_override"] = (
        float(params["nu_override"]) if params.get("nu_override") else None
    )
    return params


def interpret_results(query: str, params: dict, metrics: dict, n_paths: int, client) -> str:
    """
    Call Claude to write a plain-English risk committee briefing note.
    """
    cal_end = params.get("calibration_end") or "present"
    prompt  = INTERPRETATION_USER.format(
        scenario_name    = params["scenario_name"],
        query            = query,
        calibration_start= params["calibration_start"],
        calibration_end  = cal_end,
        horizon          = params["horizon"],
        n_paths          = n_paths,
        ann_vol          = metrics["ann_vol_pct"],
        var95            = metrics["var95_daily_pct"],
        var99            = metrics["var99_daily_pct"],
        es99             = metrics["es99_daily_pct"],
        cum_p01          = metrics["cum_p01"],
        cum_p10          = metrics["cum_p10"],
        cum_p50          = metrics["cum_p50"],
        cum_p90          = metrics["cum_p90"],
        mean_dd          = metrics["mean_dd_pct"],
        p95_dd           = metrics["p95_dd_pct"],
        prob_dd_gt20     = metrics["prob_dd_gt20_pct"],
        prob_dd_gt40     = metrics["prob_dd_gt40_pct"],
        prob_loss_gt30   = metrics["prob_loss_gt30_pct"],
    )
    response = client.messages.create(
        model="claude-opus-4-6",
        max_tokens=800,
        system=INTERPRETATION_SYSTEM,
        messages=[{"role": "user", "content": prompt}],
    )
    return response.content[0].text.strip()


# ═══════════════════════════════════════════════════════════════════════════════
# SIMULATION (same core as scenarios.py)
# ═══════════════════════════════════════════════════════════════════════════════

def _simulate(result, params: dict, n_paths: int, seed: int) -> np.ndarray:
    model_params = result.params.copy()

    if params["gamma_multiplier"] != 1.0 and "gamma[1]" in model_params:
        model_params["gamma[1]"] = float(model_params["gamma[1]"]) * params["gamma_multiplier"]
    if params["nu_override"] is not None:
        model_params["nu"] = float(params["nu_override"])
    if params["mu_shock_pct"] != 0.0:
        model_params["mu"] = float(model_params.get("mu", 0.0)) + params["mu_shock_pct"]

    iv = float(params["initial_vol_pct"] ** 2) if params["initial_vol_pct"] else None
    burn = 0 if iv is not None else 500
    horizon = params["horizon"]

    master_rng = np.random.default_rng(seed)
    path_seeds = master_rng.integers(0, 2**31, size=n_paths)

    paths = np.empty((n_paths, horizon), dtype=float)
    for i, s in enumerate(path_seeds):
        result.model.distribution._generator = np.random.default_rng(int(s))
        sim = result.model.simulate(
            model_params, nobs=horizon, burn=burn, initial_value_vol=iv
        )
        paths[i] = sim["data"].values
    return paths


def _max_drawdown(price_path: np.ndarray) -> float:
    roll = np.maximum.accumulate(price_path)
    return float(np.max((roll - price_path) / roll))


def compute_metrics(paths: np.ndarray) -> dict:
    daily = paths.ravel()
    var95 = float(-np.percentile(daily, 5))
    var99 = float(-np.percentile(daily, 1))
    tail  = daily[daily <= np.percentile(daily, 1)]
    es99  = float(-np.mean(tail)) if len(tail) else var99

    price_paths = 100.0 * np.exp(np.cumsum(paths / 100.0, axis=1))
    total_ret   = (price_paths[:, -1] - 100.0) / 100.0
    drawdowns   = np.array([_max_drawdown(price_paths[i]) for i in range(len(paths))])

    return {
        "ann_vol_pct":        round(float(np.std(daily, ddof=1) * np.sqrt(252)), 1),
        "var95_daily_pct":    round(var95, 2),
        "var99_daily_pct":    round(var99, 2),
        "es99_daily_pct":     round(es99, 2),
        "cum_p01":            round(float(np.percentile(total_ret,  1)) * 100, 1),
        "cum_p10":            round(float(np.percentile(total_ret, 10)) * 100, 1),
        "cum_p50":            round(float(np.percentile(total_ret, 50)) * 100, 1),
        "cum_p90":            round(float(np.percentile(total_ret, 90)) * 100, 1),
        "mean_dd_pct":        round(float(np.mean(drawdowns))           * 100, 1),
        "p95_dd_pct":         round(float(np.percentile(drawdowns, 95)) * 100, 1),
        "prob_dd_gt20_pct":   round(float(np.mean(drawdowns > 0.20))    * 100, 1),
        "prob_dd_gt40_pct":   round(float(np.mean(drawdowns > 0.40))    * 100, 1),
        "prob_loss_gt30_pct": round(float(np.mean(total_ret < -0.30))   * 100, 1),
        "price_paths":        price_paths,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# FAN CHART
# ═══════════════════════════════════════════════════════════════════════════════

def make_fan_chart(price_paths: np.ndarray, scenario_name: str) -> bytes:
    horizon = price_paths.shape[1]
    days    = np.arange(horizon)
    pcts    = [1, 10, 25, 50, 75, 90, 99]
    bands   = np.percentile(price_paths, pcts, axis=0)
    color   = "#1a365d"

    fig, ax = plt.subplots(figsize=(7, 4))
    ax.fill_between(days, bands[0], bands[6], alpha=0.08, color=color)
    ax.fill_between(days, bands[1], bands[5], alpha=0.15, color=color)
    ax.fill_between(days, bands[2], bands[4], alpha=0.25, color=color, label="P25–P75")
    ax.plot(days, bands[3], color=color, lw=2.0, label="Median")
    ax.plot(days, bands[0], color=color, lw=0.7, ls="--", alpha=0.5, label="P1 / P99")
    ax.plot(days, bands[6], color=color, lw=0.7, ls="--", alpha=0.5)
    ax.axhline(100, color="#718096", lw=0.8, ls=":", alpha=0.7, label="Start (100)")

    ax.set_title(scenario_name, fontsize=10, fontweight="bold", color="#1a365d", pad=8)
    ax.set_xlabel("Trading days", fontsize=9)
    ax.set_ylabel("Price index (start = 100)", fontsize=9)
    ax.tick_params(labelsize=8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(fontsize=8, loc="upper left", framealpha=0.7)
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


# ═══════════════════════════════════════════════════════════════════════════════
# WORD DOC OUTPUT
# ═══════════════════════════════════════════════════════════════════════════════

def save_word_doc(
    query: str,
    params: dict,
    metrics: dict,
    briefing: str,
    fan_png: bytes,
    n_paths: int,
    output_path: Path,
) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def _rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def _shade(cell, h):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  h.lstrip("#"))
        tcPr.append(shd)

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin  = sec.bottom_margin = Inches(0.9)

    # ── header ───────────────────────────────────────────────────────────────
    h = doc.add_heading(params["scenario_name"], level=1)
    for r in h.runs:
        r.font.color.rgb = _rgb("#1a365d")
        r.font.size      = Pt(16)

    p = doc.add_paragraph()
    r = p.add_run(f"Generated: {date.today().strftime('%d %B %Y')}  ·  "
                  f"Model: GJR-GARCH(1,1) + Student-t  ·  Paths: {n_paths:,}")
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb("#718096")
    r.italic = True

    # ── analyst query box ────────────────────────────────────────────────────
    doc.add_paragraph()
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    c = tbl.rows[0].cells[0]
    _shade(c, "#ebf8ff")
    c.text = ""
    rp = c.paragraphs[0]
    rb = rp.add_run("Analyst query:  ")
    rb.bold      = True
    rb.font.size = Pt(9)
    rn = rp.add_run(f'"{query}"')
    rn.font.size = Pt(9)
    rn.italic    = True

    doc.add_paragraph()

    # ── model assumptions ─────────────────────────────────────────────────────
    doc.add_heading("Model Assumptions", level=2)
    atbl = doc.add_table(rows=0, cols=2)
    atbl.style = "Table Grid"

    def add_row(label, value, shade="#f7fafc"):
        row = atbl.add_row()
        _shade(row.cells[0], shade)
        _shade(row.cells[1], shade)
        row.cells[0].text = label
        row.cells[1].text = value
        for c in row.cells:
            c.paragraphs[0].runs[0].font.size = Pt(8)

    cal_end = params.get("calibration_end") or "present"
    add_row("Calibration window",  f"{params['calibration_start']} → {cal_end}")
    add_row("Horizon",             f"{params['horizon']} trading days",        "#ffffff")
    add_row("Starting vol (σ₀)",   f"{params['initial_vol_pct']} % / day"
            if params['initial_vol_pct'] else "Unconditional (fitted)")
    add_row("Leverage (γ)",        f"{params['gamma_multiplier']}× fitted",    "#ffffff")
    add_row("Tail heaviness (ν)",  str(params['nu_override'])
            if params['nu_override'] else "Fitted (~5.2)")
    add_row("Mean drift shock",    f"{params['mu_shock_pct']:+.3f} % / day"
            if params['mu_shock_pct'] else "None",                             "#ffffff")

    # parameter reasoning
    doc.add_paragraph()
    rp = doc.add_paragraph()
    rb = rp.add_run("Parameter reasoning:  ")
    rb.bold = True; rb.font.size = Pt(9)
    rn = rp.add_run(params.get("reasoning", ""))
    rn.font.size = Pt(9); rn.font.color.rgb = _rgb("#2d3748")

    doc.add_paragraph()

    # ── risk metrics ──────────────────────────────────────────────────────────
    doc.add_heading("Risk Metrics", level=2)
    mtbl = doc.add_table(rows=0, cols=2)
    mtbl.style = "Table Grid"

    # header
    hr = mtbl.add_row()
    for c, txt in zip(hr.cells, ["Metric", "Value"]):
        _shade(c, "#1a365d")
        c.text = txt
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size  = Pt(8)
        c.paragraphs[0].runs[0].font.color.rgb = _rgb("#ffffff")

    rows = [
        ("Annualised volatility",      f"{metrics['ann_vol_pct']:.1f} %"),
        ("Daily VaR 95%",              f"{metrics['var95_daily_pct']:.2f} %"),
        ("Daily VaR 99%",              f"{metrics['var99_daily_pct']:.2f} %"),
        ("Daily ES 99%",               f"{metrics['es99_daily_pct']:.2f} %"),
        ("Horizon return — worst 1%",  f"{metrics['cum_p01']:.1f} %"),
        ("Horizon return — worst 10%", f"{metrics['cum_p10']:.1f} %"),
        ("Horizon return — median",    f"{metrics['cum_p50']:.1f} %"),
        ("Mean max drawdown",          f"{metrics['mean_dd_pct']:.1f} %"),
        ("P95 max drawdown",           f"{metrics['p95_dd_pct']:.1f} %"),
        ("Prob(drawdown > 20 %)",      f"{metrics['prob_dd_gt20_pct']:.1f} %"),
        ("Prob(drawdown > 40 %)",      f"{metrics['prob_dd_gt40_pct']:.1f} %"),
        ("Prob(total loss > 30 %)",    f"{metrics['prob_loss_gt30_pct']:.1f} %"),
    ]
    for i, (label, value) in enumerate(rows):
        row = mtbl.add_row()
        shade = "#f7fafc" if i % 2 == 0 else "#ffffff"
        _shade(row.cells[0], shade); _shade(row.cells[1], shade)
        row.cells[0].text = label
        row.cells[1].text = value
        for c in row.cells:
            c.paragraphs[0].runs[0].font.size = Pt(8)
        row.cells[1].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()

    # ── fan chart ─────────────────────────────────────────────────────────────
    doc.add_picture(io.BytesIO(fan_png), width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # ── risk committee briefing ───────────────────────────────────────────────
    doc.add_heading("Risk Committee Briefing Note", level=2)
    for i, para in enumerate(briefing.split("\n\n")):
        para = para.strip()
        if not para:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        # Bold the first sentence of each paragraph
        sentences = para.split(". ", 1)
        if len(sentences) == 2:
            rb = p.add_run(sentences[0] + ". ")
            rb.bold = True; rb.font.size = Pt(9)
            rn = p.add_run(sentences[1])
            rn.font.size = Pt(9)
        else:
            rn = p.add_run(para)
            rn.font.size = Pt(9)

    doc.save(output_path)


# ═══════════════════════════════════════════════════════════════════════════════
# CONSOLE OUTPUT
# ═══════════════════════════════════════════════════════════════════════════════

def _print_params(params: dict) -> None:
    cal_end = params.get("calibration_end") or "present"
    iv      = f"{params['initial_vol_pct']} % / day" if params["initial_vol_pct"] else "Unconditional"
    nu      = str(params["nu_override"]) if params["nu_override"] else "Fitted"
    drift   = f"{params['mu_shock_pct']:+.3f} % / day" if params["mu_shock_pct"] else "None"

    print(f"""
  ┌─────────────────────────────────────────────────────────────┐
  │  Extracted Parameters                                        │
  ├─────────────────────────────────────────────────────────────┤
  │  Scenario name      {params['scenario_name']:<38}│
  │  Calibration        {params['calibration_start']} → {cal_end:<25}│
  │  Horizon            {params['horizon']:<38}│
  │  Starting vol       {iv:<38}│
  │  Leverage (γ)       {params['gamma_multiplier']:<38}│
  │  Tail heaviness (ν) {nu:<38}│
  │  Drift shock        {drift:<38}│
  └─────────────────────────────────────────────────────────────┘

  Reasoning: {textwrap.fill(params['reasoning'], 70, subsequent_indent='             ')}""")


def _print_metrics(metrics: dict) -> None:
    print(f"""
  ┌─────────────────────────────────────────────────────────────┐
  │  Risk Metrics                                                │
  ├─────────────────────────────────────────────────────────────┤
  │  Annualised vol          {metrics['ann_vol_pct']:>6.1f} %                      │
  │  Daily VaR 99%           {metrics['var99_daily_pct']:>6.2f} %                      │
  │  Daily ES  99%           {metrics['es99_daily_pct']:>6.2f} %                      │
  ├─────────────────────────────────────────────────────────────┤
  │  Horizon P1  cum. return {metrics['cum_p01']:>6.1f} %                      │
  │  Horizon P10 cum. return {metrics['cum_p10']:>6.1f} %                      │
  │  Horizon median  return  {metrics['cum_p50']:>6.1f} %                      │
  ├─────────────────────────────────────────────────────────────┤
  │  Mean max drawdown       {metrics['mean_dd_pct']:>6.1f} %                      │
  │  P95  max drawdown       {metrics['p95_dd_pct']:>6.1f} %                      │
  │  Prob(drawdown > 20 %)   {metrics['prob_dd_gt20_pct']:>6.1f} %                      │
  │  Prob(drawdown > 40 %)   {metrics['prob_dd_gt40_pct']:>6.1f} %                      │
  └─────────────────────────────────────────────────────────────┘""")


def _print_briefing(briefing: str) -> None:
    divider = "  " + "━" * 65
    print(f"\n{divider}")
    print("  RISK COMMITTEE BRIEFING NOTE")
    print(f"{divider}\n")
    for para in briefing.split("\n\n"):
        para = para.strip()
        if para:
            print(textwrap.fill(para, 67, initial_indent="  ",
                                subsequent_indent="  "))
            print()
    print(divider)


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

def run_once(query: str, n_paths: int, api_key: str | None, yes: bool) -> None:
    client = _get_client(api_key)

    # ── 1. Extract parameters ────────────────────────────────────────────────
    print("\n  Interpreting scenario ...", end=" ", flush=True)
    params = extract_params(query, client)
    print("done")
    _print_params(params)

    # ── 2. Confirm ───────────────────────────────────────────────────────────
    if not yes:
        answer = input("\n  Proceed with this scenario? [Y/n]: ").strip().lower()
        if answer == "n":
            print("  Cancelled.\n")
            return

    # ── 3. Fetch data ─────────────────────────────────────────────────────────
    print(f"\n  Fetching data ...", end=" ", flush=True)
    df = fetch_brent(start=params["calibration_start"])
    if params.get("calibration_end"):
        df = df[df.index <= params["calibration_end"]]
    returns = df["log_return"].dropna()
    print(f"done  ({len(returns):,} observations)")

    # ── 4. Fit model ─────────────────────────────────────────────────────────
    print(f"  Fitting GJR-GARCH ...", end=" ", flush=True)
    result = fit_gjr_garch(returns)
    print(f"done  (ν={result.params.get('nu', float('nan')):.2f}  "
          f"β={result.params.get('beta[1]', float('nan')):.4f})")

    # ── 5. Simulate ──────────────────────────────────────────────────────────
    print(f"  Simulating {n_paths:,} paths × {params['horizon']} days ...",
          end=" ", flush=True)
    paths = _simulate(result, params, n_paths, SEED)
    print("done")

    # ── 6. Metrics ───────────────────────────────────────────────────────────
    metrics = compute_metrics(paths)
    _print_metrics(metrics)

    # ── 7. Interpret results ──────────────────────────────────────────────────
    print("\n  Generating risk committee briefing ...", end=" ", flush=True)
    briefing = interpret_results(query, params, metrics, n_paths, client)
    print("done")
    _print_briefing(briefing)

    # ── 8. Fan chart + Word doc ──────────────────────────────────────────────
    fan_png = make_fan_chart(metrics["price_paths"], params["scenario_name"])

    slug    = params["scenario_name"].lower().replace(" ", "-")[:30]
    outfile = REPORTS / f"nl-{slug}-{date.today().isoformat()}.docx"
    save_word_doc(query, params, metrics, briefing, fan_png, n_paths, outfile)
    print(f"\n  Report saved → {outfile.name}\n")


def interactive_loop(n_paths: int, api_key: str | None) -> None:
    print("\n" + "═" * 67)
    print("  Brent Crude Stress Engine — Natural Language Mode")
    print("  Type a scenario in plain English. Type 'quit' to exit.")
    print("═" * 67)

    while True:
        print()
        try:
            query = input("  Scenario > ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\n  Exiting.\n")
            break

        if not query:
            continue
        if query.lower() in {"quit", "exit", "q"}:
            print("  Exiting.\n")
            break

        try:
            run_once(query, n_paths=n_paths, api_key=api_key, yes=False)
        except Exception as exc:
            print(f"\n  ✗ Error: {exc}\n")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Natural language Brent crude stress test → Word doc"
    )
    parser.add_argument(
        "--query", "-q", type=str, default=None,
        help="Scenario description in plain English. If omitted, enters interactive mode."
    )
    parser.add_argument(
        "--n-paths", type=int, default=1000,
        help="Monte Carlo paths (default 1000)."
    )
    parser.add_argument(
        "--api-key", type=str, default=None,
        help="Anthropic API key. Defaults to ANTHROPIC_API_KEY env var."
    )
    parser.add_argument(
        "--yes", "-y", action="store_true",
        help="Skip confirmation prompt and run immediately."
    )
    args = parser.parse_args()

    if args.query:
        run_once(args.query, n_paths=args.n_paths,
                 api_key=args.api_key, yes=args.yes)
    else:
        interactive_loop(n_paths=args.n_paths, api_key=args.api_key)


if __name__ == "__main__":
    main()
